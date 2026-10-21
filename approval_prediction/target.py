"""WIN/LOST depends only on presence of the three original source cells."""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .config import ROOT

TARGET_DEFINITION = """Trong phạm vi đề tài, một đơn được gán WIN khi cả ba trường Sản phẩm, Số lượng và Đơn vị đều có dữ liệu; nếu thiếu ít nhất một trong ba trường, đơn được gán LOST.

CHOT_DON là binary target: 1 = WIN, 0 = LOST. Có dữ liệu nghĩa là giá trị không null và chuyển sang text rồi strip whitespace không rỗng. Với đơn nhiều dòng, tất cả dòng nguồn phải có đủ ba trường để đơn WIN; chỉ một dòng thiếu thì LOST. Đơn không có dòng cũng không được gán WIN.

Nhãn đọc ô Excel gốc theo source_file/source_sheet/source_excel_row, trước chuyển kiểu, mapping hoặc quy đổi. Giá trị như KM, M .), quantity 0, chuỗi không phải số hoặc unit ngoài danh sách vẫn được tính là có dữ liệu. Data quality được audit riêng và không quyết định target. product_present, quantity_present, unit_present, line_complete và mọi cờ tạo nhãn đều bị loại khỏi predictors. Rule thực thi tập trung tại approval_prediction/target.py."""
SOURCE_FIELDS = {'SẢN PHẨM': 'raw_product', 'SỐ LƯỢNG': 'raw_quantity', 'ĐƠN VỊ': 'raw_unit'}


def is_present(value):
    """Null/blank only: do not parse numbers, classify products or inspect units."""
    return not pd.isna(value) and bool(str(value).strip())


def read_source_fields(items, raw_dir=None):
    """Attach original cell values by physical lineage, without pandas NA inference."""
    raw_dir = Path(raw_dir) if raw_dir is not None else ROOT/'data/raw'
    root = raw_dir.resolve()
    lineage = ['source_file', 'source_sheet', 'source_excel_row']
    if items[lineage].isna().any().any() or items.duplicated(lineage).any():
        raise ValueError('Source item lineage must be complete and unique')
    records = []
    for (filename, sheet), group in items.groupby(['source_file', 'source_sheet'], sort=False):
        path = (root/filename).resolve()
        if path.parent != root:
            raise ValueError('Source workbook must be directly inside the raw directory')
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet]
            header = next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True))
            if any(header.count(name) != 1 for name in SOURCE_FIELDS):
                raise ValueError('Source target columns must exist exactly once')
            columns = {header.index(name): target for name, target in SOURCE_FIELDS.items()}
            wanted = set(group.source_excel_row)
            if any(int(row) != row or row < 3 or row > worksheet.max_row for row in wanted):
                raise ValueError('Source Excel row must reference an existing data row')
            for row_number, values in enumerate(worksheet.iter_rows(min_row=3, max_row=int(max(wanted)), values_only=True), 3):
                if row_number in wanted:
                    records.append(dict(source_file=filename, source_sheet=sheet, source_excel_row=row_number,
                        **{target: values[index] for index, target in columns.items()}))
        finally:
            workbook.close()
    raw = pd.DataFrame(records, columns=lineage + list(SOURCE_FIELDS.values()))
    result = items.merge(raw, on=lineage, how='left', validate='one_to_one', indicator=True)
    if not result['_merge'].eq('both').all():
        raise ValueError('A source item could not be located in the workbook')
    return result.drop(columns='_merge')


def item_labels(items):
    """Input must expose raw_product/raw_quantity/raw_unit, not cleaned substitutes."""
    d = items.copy()
    for component in ['product', 'quantity', 'unit']:
        d[f'{component}_present'] = d[f'raw_{component}'].map(is_present).astype('int64')
    flags = d[['product_present', 'quantity_present', 'unit_present']]
    d['line_complete'] = flags.eq(1).all(axis=1).astype('int64')
    d['item_target'] = d.line_complete
    d['incomplete_fields'] = flags.apply(lambda row: ','.join(
        name.removesuffix('_present') for name, value in row.items() if value == 0), axis=1)
    return d


def aggregate_target(values):
    """All required source lines must have three present cells; no lines is LOST."""
    return int(len(values) > 0 and values.eq(1).all())
